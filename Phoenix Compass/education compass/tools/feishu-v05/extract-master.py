#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
从 V0.5.1 母版抽取飞书建表合同。

母版是唯一基线，这个脚本不做任何字段发明：表名、列名、列序、下拉白名单
全部照抄，只补一件事 —— 把每列映射成飞书多维表格字段类型。

V0.5.1 由 build-master-v051.py 按 feishuv0.5.docx 从 9/11 底稿生成；
底稿 Phoenix_Feishu_Operating_Model_V0.5_Clean_Master.xlsx 留作历史参照，
需要时可显式传参抽取它。

用法：
    python tools/feishu-v05/extract-master.py
    python tools/feishu-v05/extract-master.py <master.xlsx> <out.json>
"""

import io
import json
import os
import re
import sys

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_XLSX = os.path.join(HERE, 'Phoenix_Feishu_Operating_Model_V0.5.1_Master.xlsx')
DEFAULT_JSON = os.path.join(HERE, 'master-contract.json')

# 飞书多维表格字段类型
TEXT, NUMBER, SELECT, CHECKBOX = 1, 2, 3, 7

# 按运营镜像手册的约定：时间一律走单行文本，不建飞书日期字段
TIME_SUFFIXES = (' At', ' Date', ' From', ' To')
# "1.0" 之类的版本号必须留作文本，否则会被当成数字
TEXT_FORCED = ('Version',)


def parse_options(formula1):
    """把 data validation 的 formula1 拆成选项列表。"""
    if not formula1:
        return None
    value = formula1.strip()
    if value.startswith('"') and value.endswith('"'):
        value = value[1:-1]
    elif value.startswith('='):
        # 引用其他区域的下拉，这里不展开
        return None
    options = [item.strip() for item in value.split(',') if item.strip()]
    return options or None


def dropdown_by_column(worksheet):
    """列字母 -> 下拉选项。"""
    result = {}
    for validation in worksheet.data_validations.dataValidation:
        if validation.type != 'list':
            continue
        options = parse_options(validation.formula1)
        if not options:
            continue
        for cell_range in validation.sqref.ranges:
            for index in range(cell_range.min_col, cell_range.max_col + 1):
                result[get_column_letter(index)] = options
    return result


def infer_type(header, sample, options):
    if options:
        return SELECT
    if isinstance(sample, bool):
        return CHECKBOX
    if header.endswith(TIME_SUFFIXES):
        return TEXT
    if header.endswith(TEXT_FORCED):
        return TEXT
    if isinstance(sample, (int, float)):
        return NUMBER
    return TEXT


def first_data_row(worksheet, width):
    """母版每张表带一行示例数据；空行只有复选框的 False，要跳过。"""
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        cells = list(row)[:width]
        if any(cell not in (None, '', False) for cell in cells):
            return cells
    return [None] * width


def extract(xlsx_path):
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheets = []
    readme = []

    for name in workbook.sheetnames:
        worksheet = workbook[name]
        if name == 'README':
            for row in worksheet.iter_rows(values_only=True):
                cells = ['' if cell is None else str(cell) for cell in row]
                while cells and cells[-1] == '':
                    cells.pop()
                if cells:
                    readme.append(cells)
            continue

        headers = []
        for cell in next(worksheet.iter_rows(min_row=1, max_row=1)):
            headers.append(cell.value)
        while headers and headers[-1] in (None, ''):
            headers.pop()
        headers = [str(header).strip() for header in headers]

        dropdowns = dropdown_by_column(worksheet)
        sample = first_data_row(worksheet, len(headers))

        fields = []
        for index, header in enumerate(headers):
            letter = get_column_letter(index + 1)
            options = dropdowns.get(letter)
            value = sample[index] if index < len(sample) else None
            field = {
                'name': header,
                'column': letter,
                'type': infer_type(header, value, options),
                'primary': index == 0,
            }
            if options:
                field['options'] = options
            if value is not None and value != '':
                field['sample'] = value
            fields.append(field)

        sheets.append({
            'sheet': name,
            'unique': headers[0] if headers else None,
            'fields': fields,
        })

    return {
        'source_workbook': os.path.basename(xlsx_path),
        'generated_by': 'tools/feishu-v05/extract-master.py',
        'readme': readme,
        'tables': sheets,
    }


def main():
    xlsx_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    json_path = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_JSON
    contract = extract(xlsx_path)
    with io.open(json_path, 'w', encoding='utf-8') as handle:
        handle.write(json.dumps(contract, ensure_ascii=False, indent=2))
        handle.write('\n')
    total_fields = sum(len(table['fields']) for table in contract['tables'])
    print('tables: %d, fields: %d -> %s' % (len(contract['tables']), total_fields, json_path))
    for table in contract['tables']:
        selects = sum(1 for field in table['fields'] if 'options' in field)
        print('  %-22s fields=%-3d dropdowns=%-2d primary=%s'
              % (table['sheet'], len(table['fields']), selects, table['unique']))


if __name__ == '__main__':
    main()
